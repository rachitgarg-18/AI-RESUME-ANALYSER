import os
import re
import PyPDF2
import docx

def extract_text_from_pdf(pdf_path: str) -> str:
    """Extract raw text from a PDF file using PyPDF2."""
    text = ""
    try:
        with open(pdf_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        print(f"Error extracting PDF text: {e}")
    return text.strip()

def extract_text_from_docx(docx_path: str) -> str:
    """Extract raw text from a DOCX file using python-docx."""
    text_chunks = []
    try:
        doc = docx.Document(docx_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text_chunks.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_chunks.append(" | ".join(row_text))
    except Exception as e:
        print(f"Error extracting DOCX text: {e}")
    return "\n".join(text_chunks).strip()

def extract_text_from_txt(txt_path: str) -> str:
    """Extract raw text from a plain TXT file."""
    try:
        with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    except Exception as e:
        print(f"Error reading TXT file: {e}")
        return ""

def parse_resume_file(file_path: str) -> str:
    """Parse text from supported resume formats."""
    _, ext = os.path.splitext(file_path.lower())
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md"]:
        return extract_text_from_txt(file_path)
    else:
        # Fallback to reading as text
        return extract_text_from_txt(file_path)

def extract_emails(text: str) -> list[str]:
    """Find email addresses in text."""
    pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    matches = re.findall(pattern, text)
    return list(set(matches))

def extract_phone_numbers(text: str) -> list[str]:
    """Find telephone numbers in text."""
    pattern = r'(?:(?:\+|00)?\d{1,3}[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}'
    matches = re.findall(pattern, text)
    valid_phones = [p.strip() for p in matches if len(re.sub(r'\D', '', p)) >= 10]
    return list(set(valid_phones))

def extract_links(text: str) -> dict:
    """Extract LinkedIn, GitHub, and Portfolio URLs."""
    links = {
        "linkedin": None,
        "github": None,
        "portfolio": []
    }
    linkedin_match = re.search(r'(https?://(?:www\.)?linkedin\.com/in/[\w\-]+)', text, re.IGNORECASE)
    if linkedin_match:
        links["linkedin"] = linkedin_match.group(1)
        
    github_match = re.search(r'(https?://(?:www\.)?github\.com/[\w\-]+)', text, re.IGNORECASE)
    if github_match:
        links["github"] = github_match.group(1)

    url_pattern = r'https?://[^\s<>"\']+'
    all_urls = re.findall(url_pattern, text)
    for u in all_urls:
        if "linkedin.com" not in u and "github.com" not in u:
            links["portfolio"].append(u)
            
    return links

def extract_candidate_name(text: str) -> str:
    """Heuristic extraction for candidate name from header lines."""
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    for line in lines[:5]:
        # Avoid common section headers or emails/phones
        if "@" in line or "http" in line or any(k in line.lower() for k in ["resume", "curriculum vitae", "profile", "contact"]):
            continue
        words = line.split()
        if 1 <= len(words) <= 4 and all(w.isalpha() or "." in w for w in words):
            return line
    return "Candidate"

def extract_sections(text: str) -> dict:
    """Detect key sections in the resume."""
    sections = {
        "summary": False,
        "skills": False,
        "experience": False,
        "education": False,
        "projects": False,
        "certifications": False
    }
    lower_text = text.lower()
    
    if any(k in lower_text for k in ["summary", "profile", "objective", "about me"]):
        sections["summary"] = True
    if any(k in lower_text for k in ["skills", "technical skills", "technologies", "competencies", "tools"]):
        sections["skills"] = True
    if any(k in lower_text for k in ["experience", "work experience", "employment", "internship", "work history"]):
        sections["experience"] = True
    if any(k in lower_text for k in ["education", "academic", "degree", "university", "college", "bachelor", "master"]):
        sections["education"] = True
    if any(k in lower_text for k in ["projects", "personal projects", "academic projects", "key projects"]):
        sections["projects"] = True
    if any(k in lower_text for k in ["certifications", "certificates", "licenses", "courses", "achievements"]):
        sections["certifications"] = True
        
    return sections
